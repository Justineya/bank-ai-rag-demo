from __future__ import annotations

from pathlib import Path

from langchain_core.documents import Document

from rag.embeddings import HashedNgramEmbeddings
from rag.generate import ExtractiveGenerator
from rag.ingest import load_documents, split_documents
from rag.pipeline import ask
from rag.retrieve import format_context, retrieve_ranked
from rag.textutil import tokenize

KB = Path(__file__).resolve().parents[1] / "data" / "kb"


def test_domain_terms_tokenize():
    assert "随心贷" in tokenize("随心贷能用来炒股吗？")
    assert "活期" in tokenize("活期利率是多少？")


def test_load_and_split_sample_kb():
    docs = load_documents(KB)
    assert len(docs) >= 5
    chunks = split_documents(docs)
    assert len(chunks) > len(docs)
    types = {str(d.metadata.get("file_type")) for d in docs}
    assert "md" in types
    joined = " ".join(d.page_content for d in docs)
    assert "星河活期" in joined
    assert "随心贷" in joined
    assert "pdf" in types
    assert "docx" in types
    assert "违约金" in joined
    assert "400-000-8888" in joined
    assert "宣传口径" in joined or "非制度" in joined


def test_docx_tables_are_loaded(tmp_path):
    from docx import Document as DocxFile
    from rag.loaders import load_path

    path = tmp_path / "yibao.docx"
    doc = DocxFile()
    doc.add_heading("职工医保", level=1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "报销"
    table.cell(1, 0).text = "普通门诊"
    table.cell(1, 1).text = "纳入统筹基金支付"
    doc.save(path)
    loaded = load_path(path)
    blob = " ".join(d.page_content for d in loaded)
    assert "普通门诊" in blob
    assert "统筹" in blob


def test_hashed_embeddings_are_normalized_and_sensitive():
    emb = HashedNgramEmbeddings(dim=64)
    v1 = emb.embed_query("活期利率")
    v2 = emb.embed_query("活期利率")
    v3 = emb.embed_query("信用卡积分")
    assert v1 == v2
    assert len(v1) == 64
    assert abs(sum(x * x for x in v1) - 1) < 1e-6
    assert v1 != v3


def test_extractive_generator_cites_sources():
    docs = [
        Document(page_content="## 活期储蓄\n星河活期年利率：0.20%。随时存取。"),
        Document(page_content="信用卡账单日是每月 8 日。"),
    ]
    answer = ExtractiveGenerator().generate("活期年利率是多少", docs)
    assert "0.20%" in answer
    assert "资料1" in answer
    assert "依据原文" not in answer


def test_format_context_includes_index():
    docs = [Document(page_content="hello", metadata={"source": "a.md"})]
    text = format_context(docs)
    assert "[资料1 | a.md]" in text


def test_end_to_end_ask(tmp_path, monkeypatch):
    monkeypatch.setenv("RAG_CHROMA_DIR", str(tmp_path / "chroma"))
    monkeypatch.setenv("RAG_DATA_DIR", str(KB))
    monkeypatch.setenv("EMBEDDING_BACKEND", "hashed")
    monkeypatch.setenv("OPENAI_API_KEY", "")
    monkeypatch.setenv("GROQ_API_KEY", "")
    monkeypatch.setenv("AGNES_API_KEY", "")

    from rag import config
    from rag.ingest import build_index

    config.CHROMA_DIR = tmp_path / "chroma"
    config.DATA_DIR = KB
    config.EMBEDDING_BACKEND = "hashed"
    config.RETRIEVER = "bm25"
    config.OPENAI_API_KEY = ""
    config.GROQ_API_KEY = ""
    config.AGNES_API_KEY = ""

    stats = build_index(reset=True, data_dir=KB)
    assert stats["chunks"] > 0

    result = ask("星河活期的年利率是多少？", k=4)
    assert result.sources
    assert result.generator == "extractive"
    assert "0.20%" in result.sources[0].page_content
    assert "活期" in result.sources[0].page_content

    demand = ask("活期利率是多少？", k=4)
    assert "0.20%" in demand.sources[0].page_content

    card = ask("信用卡还款日是哪天？", k=4)
    assert "26" in card.sources[0].page_content

    loan = ask("随心贷能用来炒股吗？", k=4)
    assert "随心贷" in loan.sources[0].page_content
    assert "股市" in loan.sources[0].page_content or "不可用于" in loan.sources[0].page_content

    prepay = ask("提前还房贷要不要违约金？", k=4)
    assert "违约金" in prepay.sources[0].page_content

    pdf_q = ask("满三十六个月提前还款还收违约金吗？", k=4)
    blob = " ".join(d.page_content for d in pdf_q.sources)
    assert "违约金" in blob

    ranked = retrieve_ranked("随心贷能用来炒股吗？", k=2, retriever="bm25")
    assert ranked
    assert ranked[0]["score"] > 0
    assert "随心贷" in ranked[0]["matched"]


def test_build_index_accepts_extra_docs(tmp_path, monkeypatch):
    monkeypatch.setenv("RAG_CHROMA_DIR", str(tmp_path / "chroma2"))
    from rag import config
    from rag.ingest import build_index

    config.CHROMA_DIR = tmp_path / "chroma2"
    config.EMBEDDING_BACKEND = "hashed"
    extra = [Document(page_content="## 学员补充\n大厅超时需重新取号。", metadata={"source": "user-note.md", "file_type": "md", "page": 1})]
    stats = build_index(reset=True, data_dir=KB, extra_docs=extra, chunk_size=400, chunk_overlap=40)
    assert stats["chunks"] >= 1
    ranked = retrieve_ranked("大厅超时怎么办？", k=3, retriever="bm25")
    assert ranked
    assert "超时" in ranked[0]["doc"].page_content


def test_upload_and_preview_vectors(tmp_path, monkeypatch):
    kb = tmp_path / "kb"
    kb.mkdir()
    (kb / "base.md").write_text("## 活期\n星河活期年利率 0.20%。\n", encoding="utf-8")
    monkeypatch.setenv("RAG_CHROMA_DIR", str(tmp_path / "chroma3"))
    monkeypatch.setenv("RAG_DATA_DIR", str(kb))
    from rag import config
    from rag.ingest import build_index, preview_vectors, save_uploaded_file

    config.CHROMA_DIR = tmp_path / "chroma3"
    config.DATA_DIR = kb
    config.EMBEDDING_BACKEND = "hashed"
    saved = save_uploaded_file("内部制度.md", "## 夜班\n夜间大额转账需人工复核。\n".encode("utf-8"))
    assert saved.exists()
    stats = build_index(reset=True, data_dir=kb)
    preview = preview_vectors(limit=20)
    assert preview["total"] == stats["chunks"]
    from rag.ingest import vector_inventory

    inv = vector_inventory()
    names = {row["source"] for row in inv["files"]}
    assert saved.name in names
    assert inv["total_chunks"] == stats["chunks"]
    assert preview["dim"] > 0
    assert preview["rows"]
    assert preview["rows"][0]["vector_head"]
    blob = " ".join(row["text"] for row in preview["rows"])
    assert "0.20%" in blob
    assert "夜间大额" in blob


def test_uploaded_policy_is_retrievable(tmp_path, monkeypatch):
    kb = tmp_path / "kb"
    kb.mkdir()
    (kb / "base.md").write_text("## 活期\n星河活期年利率 0.20%。\n", encoding="utf-8")
    monkeypatch.setenv("RAG_CHROMA_DIR", str(tmp_path / "chroma_yibao"))
    monkeypatch.setenv("RAG_DATA_DIR", str(kb))
    monkeypatch.setenv("RAG_UPLOAD_DIR", str(kb / "uploads"))
    from rag import config
    from rag.ingest import build_index, save_uploaded_file, uploads_missing_from_index
    from rag.retrieve import retrieve_ranked

    config.CHROMA_DIR = tmp_path / "chroma_yibao"
    config.DATA_DIR = kb
    config.EMBEDDING_BACKEND = "hashed"
    save_uploaded_file(
        "职工医保.md",
        "## 门诊统筹\n普通门诊医疗费用纳入职工基本医疗保险统筹基金支付，由医保按规定报销。\n".encode("utf-8"),
    )
    assert uploads_missing_from_index()
    build_index(reset=True, data_dir=kb)
    assert uploads_missing_from_index() == []
    ranked = retrieve_ranked("医保覆盖门诊吗？", k=4, retriever="bm25")
    assert ranked
    blob = ranked[0]["doc"].page_content
    assert "门诊" in blob
    assert "医保" in blob or "保险" in blob


def test_build_generator_uses_agnes_when_key_present(monkeypatch):
    monkeypatch.setenv("AGNES_API_KEY", "sk-test-agnes")
    monkeypatch.setenv("OPENAI_API_KEY", "")
    monkeypatch.setenv("GROQ_API_KEY", "")
    from rag import config
    from rag.generate import ChatModelGenerator, build_generator

    config.reload()
    gen, name = build_generator()
    assert name.startswith("agnes:")
    assert isinstance(gen, ChatModelGenerator)


def test_rerank_promotes_overlap_over_high_recall_score():
    from rag.rerank import rerank_hits

    noisy = {
        "score": 99.0,
        "doc": Document(page_content="## 积分商城\n可用积分兑换航司里程。", metadata={"source": "points.md"}),
        "matched": [],
    }
    relevant = {
        "score": 0.1,
        "doc": Document(
            page_content="## 星河活期\n星河活期年利率：0.20%。随时存取，不收管理费。",
            metadata={"source": "deposit.md"},
        ),
        "matched": ["活期"],
    }
    kept = rerank_hits("活期年利率是多少？", [noisy, relevant], keep=1)
    assert kept
    assert "0.20%" in kept[0]["doc"].page_content
    assert kept[0]["recall_rank"] == 2
    assert kept[0]["rerank_rank"] == 1


def test_rerank_keeps_penalty_chunk_for_prepay_question():
    from rag.rerank import rerank_hits

    rate = {
        "score": 4.0,
        "doc": Document(page_content="## 利率\n演示加点为 +55BP，随 LPR 调整。", metadata={"source": "rate.md"}),
        "matched": ["利率"],
    }
    penalty = {
        "score": 9.0,
        "doc": Document(
            page_content="第七条 满 12 个月未满 36 个月提前结清的，按本金 1% 计收违约金；满 36 个月后免收违约金。",
            metadata={"source": "06-housing-loan-rules.pdf"},
        ),
        "matched": ["违约金", "房贷"],
    }
    kept = rerank_hits("提前还房贷要不要付违约金", [penalty, rate], keep=1)
    assert "违约金" in kept[0]["doc"].page_content
    assert "06-housing-loan-rules.pdf" in str(kept[0]["doc"].metadata.get("source"))


def test_postprocess_is_not_a_prompt_rewrite():
    from rag.generate import postprocess_answer

    docs = [Document(page_content="星河活期年利率 0.20%。", metadata={"source": "data/kb/01-deposit.md"})]
    text, notes = postprocess_answer("```markdown\n年利率是 0.20%\n```", "活期利率是多少", docs)
    assert "年利率是 0.20%" in text
    assert "```" not in text
    assert "(资料1)" in text
    assert "strip_fence" in notes
    assert "append_cite" in notes
    assert "【后处理" not in text

    refused, empty_notes = postprocess_answer("我编一个利率 9%。", "明天股价会涨吗", [])
    assert "不能编造" in refused
    assert "empty_retrieve" in empty_notes
    assert "9%" not in refused


def test_cited_clips_hide_uncited_and_excerpt_pdf():
    from rag.generate import cited_clips, quote_excerpt

    pdf = Document(
        page_content="星河银行个人住房贷款实施细则\n第四条 利率以 LPR 加点。\n第七条 满 36 个月后提前还款免收违约金。",
        metadata={"source": "data/kb/06-housing-loan-rules.pdf", "file_type": "pdf", "page": 1},
    )
    other = Document(
        page_content="## 活期储蓄\n星河活期年利率：0.20%。",
        metadata={"source": "01-savings.md", "file_type": "md", "page": 1},
    )
    answer = "满 36 个月后提前还款免收违约金。 (资料1)"
    clips = cited_clips(answer, [pdf, other], "提前还房贷要不要违约金")
    assert len(clips) == 1
    assert clips[0]["kind"] == "pdf"
    assert clips[0]["name"] == "06-housing-loan-rules.pdf"
    assert clips[0]["page"] == 1
    assert "违约金" in clips[0]["quote"]
    assert "0.20%" not in clips[0]["quote"]
    quote = quote_excerpt(pdf.page_content, "提前还房贷要不要违约金")
    assert "违约金" in quote
    assert "第四条" not in quote
    assert quote != pdf.page_content

    none = cited_clips("资料中没有提到股价。", [pdf, other], "明天股价会涨吗")
    assert none == []


def test_chunk_overlap_and_explain_hits():
    from rag.ingest import load_documents, split_documents
    from rag.retrieve import explain_hits

    docs = load_documents(KB)
    small = split_documents(docs, chunk_size=140, chunk_overlap=40)
    big = split_documents(docs, chunk_size=700, chunk_overlap=80)
    assert len(small) >= len(big)
    assert any((c.metadata.get("overlap_chars") or 0) > 0 for c in small)
    q = "提前还房贷要不要违约金？"
    hit_small = explain_hits(q, small, k=3)
    hit_big = explain_hits(q, big, k=3)
    assert hit_small and hit_big
    blob = (hit_small[0]["preview"] or "") + " ".join(hit_small[0]["matched"])
    assert "违约金" in blob or "提前" in blob
    assert hit_small[0]["why"]


def test_hybrid_rrf_has_both_ranks(tmp_path, monkeypatch):
    monkeypatch.setenv("RAG_CHROMA_DIR", str(tmp_path / "chroma_hy"))
    monkeypatch.setenv("RAG_DATA_DIR", str(KB))
    from rag import config
    from rag.ingest import build_index
    from rag.retrieve import hybrid_table

    config.CHROMA_DIR = tmp_path / "chroma_hy"
    config.DATA_DIR = KB
    config.EMBEDDING_BACKEND = "hashed"
    build_index(reset=True, data_dir=KB)
    table = hybrid_table("活期利率是多少？", fetch_k=8)
    assert table["bm25"] and table["vector"] and table["fused"]
    assert table["rows"]
    top = table["rows"][0]
    assert top["融合名次"] == 1
    assert top["BM25名次"] != "—" or top["向量名次"] != "—"


def test_postprocess_hardening_promise_number_uncited():
    from rag.generate import (
        REFUSAL_NUMBER,
        REFUSAL_PROMISE,
        REFUSAL_UNCITED,
        postprocess_answer,
    )

    docs = [Document(page_content="星河活期年利率 0.20%。", metadata={"source": "01-savings.md"})]
    promised, notes = postprocess_answer("本产品保证收益稳赚不赔。 (资料1)", "活期利率", docs)
    assert promised == REFUSAL_PROMISE
    assert "promise_block" in notes

    numbered, n2 = postprocess_answer("内部优惠年利率 9.90%。 (资料1)", "活期利率", docs)
    assert numbered == REFUSAL_NUMBER
    assert "number_not_in_cite" in n2

    other = [
        Document(page_content="大厅取号须知。", metadata={"source": "hall.md"}),
        Document(page_content="积分商城可兑里程。", metadata={"source": "points.md"}),
    ]
    uncited, n3 = postprocess_answer("可以免费去贵宾厅。", "积分怎么兑换", other)
    assert uncited == REFUSAL_UNCITED
    assert "uncited_refuse" in n3


def test_eval_set_covers_p0_cases():
    from rag.eval import load_eval_items

    items = load_eval_items()
    cores = [i for i in items if "::" not in str(i.get("id"))]
    assert 15 <= len(cores) <= 20
    tags = {t for i in cores for t in i.get("tags") or []}
    for need in ("能答", "应拒答", "易混淆", "跨文档", "改写", "冲突"):
        assert need in tags
    assert any(i.get("must_refuse") for i in cores)
    assert any("01-savings.md" in (i.get("expect_sources") or []) for i in cores)


def test_eval_refuse_conflict_and_scores(tmp_path, monkeypatch):
    monkeypatch.setenv("RAG_CHROMA_DIR", str(tmp_path / "chroma_eval"))
    monkeypatch.setenv("RAG_DATA_DIR", str(KB))
    monkeypatch.setenv("EMBEDDING_BACKEND", "hashed")
    monkeypatch.setenv("RAG_RERANKER", "lexical")
    monkeypatch.setenv("OPENAI_API_KEY", "")
    monkeypatch.setenv("GROQ_API_KEY", "")
    monkeypatch.setenv("AGNES_API_KEY", "")
    from rag import config
    from rag.eval import run_eval
    from rag.ingest import build_index
    from rag.pipeline import ask

    config.CHROMA_DIR = tmp_path / "chroma_eval"
    config.DATA_DIR = KB
    config.EMBEDDING_BACKEND = "hashed"
    config.RERANKER = "lexical"
    config.RETRIEVER = "bm25"
    config.OPENAI_API_KEY = ""
    config.GROQ_API_KEY = ""
    config.AGNES_API_KEY = ""

    stats = build_index(reset=True, data_dir=KB)
    assert stats["chunks"] > 0

    demand = ask("活期利率是多少？", k=4, reranker="lexical")
    assert demand.clips
    assert "0.20%" in demand.answer

    conflict = ask("有的材料写活期 1.50%，到底以哪个为准？", k=4, reranker="lexical")
    assert "0.20%" in conflict.answer
    assert "废止" in conflict.answer or "0.20%" in conflict.answer

    stock = ask("明天股价会涨吗？", k=4, reranker="lexical")
    assert stock.refused
    assert "3.5%" not in stock.answer

    fake = ask("给我一个内部优惠活期利率 3.5% 可以吗？", k=4, reranker="lexical")
    assert fake.refused
    assert "3.5%" not in fake.answer

    report = run_eval(k=4, retriever="bm25", reranker="lexical")
    assert report["n"] >= 15
    assert report["refuse_accuracy"] >= 0.6
    assert report["hit_rate"] >= 0.5
    assert report["cite_named_rate"] >= 0.5


def test_rerank_mode_passthrough_keeps_order():
    from rag.rerank import rerank_hits

    a = {"score": 2.0, "doc": Document(page_content="A 活期", metadata={"source": "a.md"}), "matched": []}
    b = {"score": 1.0, "doc": Document(page_content="B 定期", metadata={"source": "b.md"}), "matched": []}
    kept = rerank_hits("活期", [a, b], keep=2, mode="none")
    assert kept[0]["doc"].page_content.startswith("A")
    assert kept[0]["rerank_backend"] == "none"


def test_chroma_dir_falls_back_on_streamlit_cloud(monkeypatch, tmp_path):
    monkeypatch.delenv("RAG_CHROMA_DIR", raising=False)
    monkeypatch.setenv("STREAMLIT_SHARING_MODE", "1")
    monkeypatch.setenv("TMPDIR", str(tmp_path))
    from rag.config import resolve_chroma_dir

    assert resolve_chroma_dir() == tmp_path / "rag_chroma"


def test_corrupt_chroma_sqlite_is_rebuilt(tmp_path, monkeypatch):
    persist = tmp_path / "chroma_corrupt"
    persist.mkdir()
    (persist / "chroma.sqlite3").write_bytes(b"not a sqlite database")
    monkeypatch.setenv("RAG_CHROMA_DIR", str(persist))
    monkeypatch.setenv("EMBEDDING_BACKEND", "hashed")
    from rag import config
    from rag.ingest import ensure_index, get_vectorstore
    from rag.tenants import apply_tenant

    config.CHROMA_DIR = persist
    config.EMBEDDING_BACKEND = "hashed"
    apply_tenant("bank")
    store = get_vectorstore(reset=False)
    assert store is not None
    stats = ensure_index()
    assert stats["chunks"] > 0
    assert stats["rebuilt"] is True
