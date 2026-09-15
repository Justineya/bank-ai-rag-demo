"""把制度类 Markdown 导出为 PDF / Word，供 Loader 解析（教学用）。"""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfWriter
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parents[1]
KB = ROOT / "data" / "kb"


def md_to_pdf(md_path: Path, pdf_path: Path) -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    title = ParagraphStyle("cn_title", parent=styles["Title"], fontName="STSong-Light", fontSize=16, leading=22)
    heading = ParagraphStyle("cn_h", parent=styles["Heading2"], fontName="STSong-Light", fontSize=12, leading=18)
    body = ParagraphStyle("cn_body", parent=styles["Normal"], fontName="STSong-Light", fontSize=10, leading=16)

    story = []
    for raw in md_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            story.append(Spacer(1, 8))
            continue
        escaped = (
            line.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        if line.startswith("# "):
            story.append(Paragraph(escaped[2:], title))
        elif line.startswith("## "):
            story.append(Paragraph(escaped[3:], heading))
        else:
            story.append(Paragraph(escaped.lstrip("- "), body))
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        title=md_path.stem,
    )
    doc.build(story)


def md_to_docx(md_path: Path, docx_path: Path) -> None:
    document = DocxDocument()
    for raw in md_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            p = document.add_heading(line[2:], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        else:
            document.add_paragraph(line.lstrip("- "))
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(docx_path))


def main() -> None:
    md_to_pdf(KB / "06-housing-loan-rules.md", KB / "06-housing-loan-rules.pdf")
    md_to_docx(KB / "07-credit-card-charter.md", KB / "07-credit-card-charter.docx")
    writer = PdfWriter()
    writer.append(str(KB / "06-housing-loan-rules.pdf"))
    assert Path(KB / "06-housing-loan-rules.pdf").stat().st_size > 0
    print("wrote", KB / "06-housing-loan-rules.pdf")
    print("wrote", KB / "07-credit-card-charter.docx")


if __name__ == "__main__":
    main()
