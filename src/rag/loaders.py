"""把 Markdown / PDF / Word / 图片读成 LangChain Document。"""

from __future__ import annotations

import re
from pathlib import Path

from langchain_core.documents import Document

_OCR_ENGINE = None


SUPPORTED_SUFFIXES = {
    ".md",
    ".txt",
    ".pdf",
    ".docx",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".bmp",
    ".tif",
    ".tiff",
}


def load_path(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        text = _read_text_file(path).strip()
        if not text:
            return []
        return [
            Document(
                page_content=text,
                metadata={"source": str(path), "file_type": suffix.lstrip("."), "page": 1},
            )
        ]
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".docx":
        return _load_docx(path)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        return _load_image(path)
    return []


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _load_pdf(path: Path) -> list[Document]:
    docs = _load_pdf_pymupdf(path)
    if docs:
        return docs
    return _load_pdf_pypdf(path)


def _load_pdf_pypdf(path: Path) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    docs: list[Document] = []
    for i, page in enumerate(reader.pages, start=1):
        text = _normalize_pdf_text(page.extract_text() or "")
        if not text:
            continue
        docs.append(
            Document(
                page_content=text,
                metadata={"source": str(path), "file_type": "pdf", "page": i},
            )
        )
    return docs


def _load_pdf_pymupdf(path: Path) -> list[Document]:
    try:
        import fitz
    except ImportError:
        return []
    docs: list[Document] = []
    with fitz.open(str(path)) as pdf:
        for i, page in enumerate(pdf, start=1):
            text = _normalize_pdf_text(page.get_text("text") or "")
            if len(text) < 40:
                ocr = _ocr_png(_page_png(page))
                if len(ocr.strip()) > len(text):
                    text = ocr.strip()
            if not text:
                continue
            docs.append(
                Document(
                    page_content=text,
                    metadata={"source": str(path), "file_type": "pdf", "page": i},
                )
            )
    return docs


def _page_png(page) -> bytes:
    import fitz

    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
    return pix.tobytes("png")


def _load_image(path: Path) -> list[Document]:
    text = _ocr_png(path.read_bytes()).strip()
    if not text:
        return []
    return [
        Document(
            page_content=text,
            metadata={"source": str(path), "file_type": path.suffix.lstrip(".").lower(), "page": 1},
        )
    ]


def _load_docx(path: Path) -> list[Document]:
    from docx import Document as DocxFile

    parsed = DocxFile(str(path))
    blocks: list[str] = []
    current: list[str] = []
    pending_title = ""

    def flush() -> None:
        nonlocal current
        if current:
            blocks.append("\n".join(current))
            current = []

    for para in parsed.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "") if para.style else ""
        is_heading = "Heading" in style or style.startswith("Title")
        if is_heading:
            flush()
            pending_title = text
            continue
        if pending_title:
            current.append(pending_title)
            pending_title = ""
        current.append(text)
    if pending_title:
        current.append(pending_title)
    for table in parsed.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells if cell.text.strip()]
            if cells:
                rows.append("　".join(cells))
        if rows:
            flush()
            current.append("\n".join(rows))
    flush()
    return [
        Document(
            page_content=block,
            metadata={"source": str(path), "file_type": "docx", "page": i},
        )
        for i, block in enumerate(blocks, start=1)
        if block.strip()
    ]


def _normalize_pdf_text(text: str) -> str:
    compact = (text or "").strip()
    compact = re.sub(r"(\d)\s*\n+\s*", r"\1", compact)
    compact = re.sub(r"[ \t]+\n", "\n", compact)
    return compact.strip()


def _ocr_png(png: bytes) -> str:
    global _OCR_ENGINE
    if not png:
        return ""
    try:
        from rapidocr import RapidOCR
    except ImportError:
        try:
            from rapidocr_onnxruntime import RapidOCR
        except ImportError:
            return ""
    if _OCR_ENGINE is None:
        _OCR_ENGINE = RapidOCR()
    result = _OCR_ENGINE(png)
    return _ocr_result_text(result)


def _ocr_result_text(result) -> str:
    if result is None:
        return ""
    lines: list[str] = []
    payload = result
    if isinstance(result, tuple):
        payload = result[0]
    txts = getattr(payload, "txts", None)
    if txts:
        return "\n".join(str(t) for t in txts if t).strip()
    if isinstance(payload, list):
        for row in payload:
            if isinstance(row, (list, tuple)) and len(row) >= 2:
                lines.append(str(row[1]))
            elif isinstance(row, dict) and row.get("text"):
                lines.append(str(row["text"]))
            elif isinstance(row, str):
                lines.append(row)
    return "\n".join(lines).strip()
